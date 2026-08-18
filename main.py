"""Таймкоды ключевых мыслей из видео -> Excel."""

import json
import math
import os
import re
import subprocess
import sys
import time
from collections import Counter
from pathlib import Path

from faster_whisper import WhisperModel
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

SOURCE_DIR = Path(os.getenv("SOURCE_DIR", "source"))
OUT_DIR = Path(os.getenv("OUT_DIR", "out"))
WORK_DIR = Path(os.getenv("WORK_DIR", "work"))

WHISPER_MODEL = os.getenv("WHISPER_MODEL", "medium")
WHISPER_DEVICE = os.getenv("WHISPER_DEVICE", "cpu")
WHISPER_COMPUTE_TYPE = os.getenv("WHISPER_COMPUTE_TYPE", "int8")
WHISPER_LANGUAGE = os.getenv("WHISPER_LANGUAGE") or None
CPU_THREADS = int(os.getenv("CPU_THREADS", "0"))  # 0 = по числу ядер
BEAM_SIZE = int(os.getenv("BEAM_SIZE", "5"))
PROGRESS_SEC = 120.0  # шаг лога прогресса по таймлайну видео

BLOCK_SEC = float(os.getenv("BLOCK_SEC", "90"))          # целевая длина смыслового блока
MIN_BLOCK_SEC = float(os.getenv("MIN_BLOCK_SEC", "25"))  # раньше неё по паузе не режем
PAUSE_SEC = float(os.getenv("PAUSE_SEC", "1.5"))         # пауза как граница мысли
MAX_IDEA_CHARS = 260
MIN_IDEA_WORDS = 3      # короче — это обрывок, а не мысль
IDEA_SENTENCES = 2

VIDEO_EXT = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".m4v", ".ts", ".mpg", ".mpeg"}

# Служебные слова выкидываем из оценки: они есть везде и ничего не различают.
STOPWORDS = set("""
и в во не что он на я с со как а то все всё она так его но да ты к у же вы за бы по только ее её мне
было вот от меня еще ещё нет о из ему теперь когда даже ну вдруг ли если уже или ни быть был него до вас
нибудь опять уж вам сказал ведь там потом себя ничего ей может они тут где есть надо ней для мы тебя их
чем была сам чтоб без будто человек чего раз тоже себе под жизнь будет ж тогда кто этот говорил того
потому этого какой совсем ним здесь этом один почти мой тем чтобы нее неё кажется сейчас были куда зачем
всех никогда сегодня можно при наконец два об другой хоть после над больше тот через эти нас про всего
них какая много разве три эту моя впрочем хорошо свою этой перед иногда лучше чуть том нельзя такой им
более всегда конечно всю между это как-то вообще просто очень такие такая таких эта эти этих оно они мы
вы наш ваш свой своя свои который которая которые то-есть типа значит вот-так давайте давай
the a an and or but if then of to in on for with as is are was were be been being this that these those
it its i you he she we they them his her our your their not no do does did doing have has had can could
would should will just so than too very there here what which who when where why how about into over
""".split())

SENT_SPLIT = re.compile(r"(?<=[.!?…])\s+")
WORD_RE = re.compile(r"[а-яёa-z0-9][а-яёa-z0-9-]*", re.IGNORECASE)


def hms(seconds: float) -> str:
    seconds = int(seconds)
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def find_video() -> Path:
    if len(sys.argv) > 1:
        path = Path(sys.argv[1])
        if not path.exists():
            sys.exit(f"Файл не найден: {path}")
        return path
    files = sorted(p for p in SOURCE_DIR.iterdir() if p.suffix.lower() in VIDEO_EXT)
    if not files:
        sys.exit(f"В {SOURCE_DIR}/ нет видеофайлов")
    return files[0]


def extract_audio(video: Path) -> Path:
    """Whisper всё равно работает с 16 кГц моно — вытаскиваем один раз, чтобы не гонять видео целиком."""
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    wav = WORK_DIR / (video.stem + ".wav")
    if wav.exists():
        print(f"[audio] уже извлечено: {wav}")
        return wav
    print(f"[audio] извлекаю дорожку из {video.name}")
    subprocess.run(
        ["ffmpeg", "-y", "-i", str(video), "-vn", "-ac", "1", "-ar", "16000",
         "-c:a", "pcm_s16le", "-loglevel", "error", "-stats", "-stats_period", "30", str(wav)],
        check=True,
    )
    return wav


def run_asr(wav: Path, device: str, compute_type: str) -> list[dict]:
    print(f"[whisper] модель={WHISPER_MODEL} device={device} compute={compute_type}")
    model = WhisperModel(WHISPER_MODEL, device=device, compute_type=compute_type,
                         cpu_threads=CPU_THREADS)
    segments, info = model.transcribe(
        str(wav),
        language=WHISPER_LANGUAGE,
        vad_filter=True,
        vad_parameters={"min_silence_duration_ms": 500},
        beam_size=BEAM_SIZE,
    )
    print(f"[whisper] язык={info.language} длительность={hms(info.duration)}")

    result = []
    started, reported = time.monotonic(), 0.0
    for seg in segments:
        text = seg.text.strip()
        if not text:
            continue
        result.append({"start": seg.start, "end": seg.end, "text": text})
        # Отдельными строками, а не \r: docker logs не отдаёт незавершённую строку.
        if seg.end - reported >= PROGRESS_SEC:
            reported = seg.end
            elapsed = time.monotonic() - started
            speed = seg.end / elapsed
            eta = (info.duration - seg.end) / speed
            print(f"[whisper] {hms(seg.end)} / {hms(info.duration)} "
                  f"({speed:.2f}x, осталось ~{hms(eta)})", flush=True)
    return result


def transcribe(wav: Path) -> list[dict]:
    """VRAM делится с рабочим столом: OOM прилетает и на загрузке, и в середине прохода."""
    # Транскрипт кэшируем под именем модели: иначе смена модели молча переиспользует чужой.
    cache = wav.with_suffix("." + re.sub(r"[^\w.-]", "_", WHISPER_MODEL) + ".segments.json")
    if cache.exists() and os.getenv("FORCE_ASR") != "1":
        print(f"[whisper] беру готовый транскрипт: {cache.name}")
        return json.loads(cache.read_text(encoding="utf-8"))

    try:
        segments = run_asr(wav, WHISPER_DEVICE, WHISPER_COMPUTE_TYPE)
    except Exception as exc:
        if WHISPER_DEVICE == "cpu":
            raise
        print(f"[whisper] {exc}\n[whisper] откатываюсь на cpu/int8")
        segments = run_asr(wav, "cpu", "int8")

    cache.write_text(json.dumps(segments, ensure_ascii=False), encoding="utf-8")
    return segments


def split_blocks(segments: list[dict]) -> list[dict]:
    """Режем поток сегментов на смысловые блоки: по длинной паузе, но не короче MIN_BLOCK_SEC."""
    blocks, current = [], []

    def flush():
        if current:
            blocks.append({
                "start": current[0]["start"],
                "end": current[-1]["end"],
                "text": " ".join(s["text"] for s in current),
                "parts": [s["text"] for s in current],
            })
            current.clear()

    for i, seg in enumerate(segments):
        current.append(seg)
        length = seg["end"] - current[0]["start"]
        pause = segments[i + 1]["start"] - seg["end"] if i + 1 < len(segments) else 0.0
        if length >= BLOCK_SEC or (length >= MIN_BLOCK_SEC and pause >= PAUSE_SEC):
            flush()
    flush()
    return blocks


def build_idf(blocks: list[dict]) -> dict[str, float]:
    df = Counter()
    for block in blocks:
        df.update(set(tokenize(block["text"])))
    n = len(blocks)
    return {word: math.log((n + 1) / (count + 0.5)) for word, count in df.items()}


def tokenize(text: str) -> list[str]:
    return [w for w in (m.group().lower() for m in WORD_RE.finditer(text))
            if len(w) > 2 and w not in STOPWORDS]


def key_idea(block: dict, idf: dict[str, float]) -> str:
    """Экстрактивно: две самые «нагруженные» смыслом фразы блока в исходном порядке."""
    text = block["text"]
    sentences = [s.strip() for s in SENT_SPLIT.split(text) if s.strip()]
    # Модели вроде turbo отдают текст без пунктуации — тогда режем по границам сегментов.
    if len(sentences) < 2:
        sentences = [p.strip() for p in block["parts"] if p.strip()]
    if not sentences:
        return trim(text)

    scored = []
    for i, sentence in enumerate(sentences):
        words = tokenize(sentence)
        if len(words) < MIN_IDEA_WORDS:
            continue
        # Без нормировки по длине: иначе наверх всплывают короткие обрывки вроде «Ну-ка посмотрим».
        score = sum(idf.get(w, 0.0) for w in set(words))
        if sentence.endswith("?"):  # вопрос почти никогда не формулирует мысль
            score *= 0.6
        scored.append((score, i, sentence))

    if not scored:
        return trim(max(sentences, key=len))

    top = sorted(scored, reverse=True)[:IDEA_SENTENCES]
    return trim(" ".join(s for _, _, s in sorted(top, key=lambda x: x[1])))


def trim(text: str) -> str:
    text = text.strip()
    if len(text) <= MAX_IDEA_CHARS:
        return text
    return text[:MAX_IDEA_CHARS].rsplit(" ", 1)[0] + "…"


def write_excel(video: Path, blocks: list[dict], segments: list[dict], idf: dict) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"{video.stem}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Ключевые мысли"
    ws.append(["№", "Таймкод", "Секунды", "Длительность", "Ключевая мысль", "Текст блока"])

    for i, block in enumerate(blocks, 1):
        ws.append([
            i,
            hms(block["start"]),
            round(block["start"], 1),
            hms(block["end"] - block["start"]),
            key_idea(block, idf),
            block["text"],
        ])

    ts = wb.create_sheet("Транскрипт")
    ts.append(["Таймкод", "Секунды", "Текст"])
    for seg in segments:
        ts.append([hms(seg["start"]), round(seg["start"], 1), seg["text"]])

    for sheet, widths in ((ws, [5, 11, 10, 13, 70, 90]), (ts, [11, 10, 110])):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col, width in zip(sheet.iter_cols(min_row=1, max_row=1), widths):
            sheet.column_dimensions[col[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(out)
    return out


def main() -> None:
    video = find_video()
    print(f"[video] {video} ({video.stat().st_size / 2**30:.1f} ГБ)")

    segments = transcribe(extract_audio(video))
    if not segments:
        sys.exit("Речь не распознана — проверь звуковую дорожку")

    blocks = split_blocks(segments)
    idf = build_idf(blocks)
    out = write_excel(video, blocks, segments, idf)
    print(f"[done] {out} — блоков: {len(blocks)}, сегментов: {len(segments)}")


if __name__ == "__main__":
    main()
